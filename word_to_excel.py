import os
import re
import threading
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from difflib import SequenceMatcher
from collections import OrderedDict

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pyexcel_ods3 import save_data


# ============================================================
# ШАПКИ ФОРМАТОВ
# ============================================================
FORMAT_32 = [
    ("Номер формата", "AUTO"),
    ("Код СФО", None),
    ("Код НОПС", 0),
    ("НОПС в единственном числе", 1),
    ("НОПС во множественном числе", None),
    ("Класс ЕКПС", 2),
    ("Номер позиции базовых характеристик", 4),
    ("Номер характеристики обязательной к заполнению", 5),
]

FORMAT_35 = [
    ("Номер формата", "AUTO"),
    ("ФНН ПС", ["фнн"]),
    ("Номер позиции", ["поз"]),
    ("Номер измерения", ["измер"]),
    ("Код параметра", ["код", "параметр"]),
    ("Код объекта", ["код", "объект"]),
    ("Наименование характеристики", ["наимен", "характер"]),
    ("Функция соответствия", ["функц", "соответств"]),
    ("Код единицы измерения", ["единиц", "измер", "код"]),
    ("Полное наименование единицы измерения", ["единиц", "измер", "наимен"]),
    ("Сокращенное наименование единицы измерения", ["сокращ", "единиц"]),
    ("Код значения", ["значен", "код"]),
    ("Полное наименование значения", ["значен", "наимен"]),
    ("Сокращенное наименование значения", ["сокращ", "значен"]),
]

FORMAT_38 = [
    ("Номер формата", "AUTO"),
    ("Код СФО", ["сфо"]),
    ("Номер позиции", ["поз"]),
    ("Тип характеристики", ["вид", "характер"]),
    ("Код параметра", ["код", "параметр"]),
    ("Код объекта", ["код", "объект"]),
    ("Падеж объекта", ["падеж"]),
    ("Наименование характеристики (параметра и объекта)", ["наимен", "характер"]),
    ("Обозначение функции соответствия", ["функц", "соответств"]),
    ("Тип измеряемой величины", ["тип", "измер"]),
    ("Код единицы измерения", ["единиц", "измер", "код"]),
    ("Полное наименование единицы измерения", ["единиц", "измер", "наимен"]),
    ("Сокращенное наименование единицы измерения", ["сокращ", "единиц"]),
    ("Код возможного значения условия", ["возможн", "услов", "код"]),
    ("Полное наименование возможного значения условия", ["возможн", "услов", "наимен"]),
    ("Сокращенное наименование возможного значения условия", ["сокращ", "возможн"]),
    ("Код характера изменения свойства", ["характер", "изменен", "код"]),
    ("Нижняя граница значения характера изменения свойства", ["нижн", "границ"]),
    ("Верхняя граница значения характера изменения свойства", ["верхн", "границ"]),
    ("Номер блока характеристик", ["блок", "характер"]),
]

FORMATS = {
    "32": FORMAT_32,
    "35": FORMAT_35,
    "38": FORMAT_38,
}

FUZZY_THRESHOLD = 0.75

# Расширения для каждого выходного формата
EXT_MAP = {
    "xlsx": ".xlsx",
    "xlsm": ".xlsm",
    "ods": ".ods",
}


# ============================================================
# ЛОГИКА
# ============================================================
def normalize(text):
    if text is None:
        return ""
    text = str(text).strip().lower().replace("\n", " ")
    text = re.sub(r"[«»\"'`]", "", text)
    text = re.sub(r"[.:;,!?()\[\]{}]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def similarity(a, b):
    a, b = normalize(a), normalize(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.95
    return SequenceMatcher(None, a, b).ratio()


def has_all_markers(text, markers):
    if not markers:
        return True
    t = normalize(text)
    return all(m in t for m in markers)


# ---------- ЧТЕНИЕ WORD ----------
def clean_cell_text(cell):
    paragraphs = [p.text.strip() for p in cell.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    seen = set()
    result = []
    for p in paragraphs:
        key = re.sub(r"[^\w]+", "", p.lower(), flags=re.UNICODE)
        if key and key not in seen:
            seen.add(key)
            result.append(p)
    return " ".join(result).strip()


def get_grid_span(tc):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return 1
    gs = tcPr.find(qn("w:gridSpan"))
    if gs is None:
        return 1
    return int(gs.get(qn("w:val"), "1"))


def get_vmerge(tc):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return None
    vm = tcPr.find(qn("w:vMerge"))
    if vm is None:
        return None
    val = vm.get(qn("w:val"))
    return val if val else "continue"


def read_word_table(table):
    n_cols = len(table.columns)
    vmerge_values = [None] * n_cols

    result = []
    for row in table.rows:
        row_values = [None] * n_cols
        col_idx = 0

        for tc in row._tr.findall(qn("w:tc")):
            if col_idx >= n_cols:
                break
            span = get_grid_span(tc)
            vmerge = get_vmerge(tc)

            if vmerge == "continue":
                value = vmerge_values[col_idx]
            else:
                cell_obj = _Cell(tc, table)
                value = clean_cell_text(cell_obj)
                if vmerge == "restart":
                    vmerge_values[col_idx] = value
                else:
                    vmerge_values[col_idx] = None

            for _ in range(span):
                if col_idx < n_cols:
                    row_values[col_idx] = value
                    col_idx += 1

        while col_idx < n_cols:
            row_values[col_idx] = vmerge_values[col_idx]
            col_idx += 1

        result.append(row_values)

    return result


def read_word_tables(path):
    doc = Document(path)
    all_rows = []
    for t in doc.tables:
        all_rows.extend(read_word_table(t))
    return all_rows


def is_number_row(row):
    if not row:
        return False
    non_empty = [c for c in row if c and str(c).strip()]
    if not non_empty:
        return False
    nums = sum(1 for c in non_empty if re.fullmatch(r"\d+", str(c).strip()))
    return nums >= len(non_empty) * 0.7


def detect_word_header(rows):
    markers = ["наимен", "код", "номер", "значен", "тип", "единица",
               "характер", "функц", "измер", "поз", "вид", "объект",
               "нопс", "екпс", "класс", "описание"]

    best_idx, best_score = 0, -1
    for i, row in enumerate(rows[:10]):
        joined = normalize(" ".join(str(c) for c in row if c))
        score = sum(1 for m in markers if m in joined)
        numeric = sum(1 for c in row if c and re.fullmatch(r"\d+", str(c).strip()))
        score -= numeric * 0.5
        if score > best_score:
            best_score, best_idx = score, i

    top = rows[best_idx]
    bottom = rows[best_idx + 1] if best_idx + 1 < len(rows) else None

    if bottom is not None and is_number_row(bottom):
        merged = list(top)
        data_start = best_idx + 2
    else:
        if bottom is not None and not any(str(c).strip() for c in bottom if c):
            merged = list(top)
            data_start = best_idx + 1
        else:
            merged = []
            n = max(len(top), len(bottom) if bottom else 0)
            for i in range(n):
                t = top[i] if i < len(top) else ""
                b = bottom[i] if bottom and i < len(bottom) else ""
                t_str = str(t).strip() if t else ""
                b_str = str(b).strip() if b else ""
                if t_str and b_str and t_str != b_str:
                    merged.append(f"{t_str} {b_str}")
                elif t_str:
                    merged.append(t_str)
                elif b_str:
                    merged.append(b_str)
                else:
                    merged.append("")
            data_start = best_idx + 2

    if data_start < len(rows) and is_number_row(rows[data_start]):
        data_start += 1

    return best_idx, merged, rows[data_start:]


def build_mapping(word_headers, fmt_columns):
    hard_mode = True
    for name, val in fmt_columns:
        if isinstance(val, list):
            hard_mode = False
            break

    mapping = []

    if hard_mode:
        for name, val in fmt_columns:
            if val == "AUTO":
                mapping.append("AUTO")
            else:
                mapping.append(val)
        return mapping

    used = set()
    for fmt_name, markers in fmt_columns:
        if markers == "AUTO":
            mapping.append("AUTO")
            continue
        if markers is None:
            mapping.append(None)
            continue

        best_idx, best_score = None, 0.0
        for w_idx, w_name in enumerate(word_headers):
            if w_idx in used:
                continue
            if not w_name:
                continue
            if not has_all_markers(w_name, markers):
                continue
            s = similarity(fmt_name, w_name)
            if s > best_score:
                best_score, best_idx = s, w_idx

        if best_idx is not None and best_score >= FUZZY_THRESHOLD:
            mapping.append(best_idx)
            used.add(best_idx)
        else:
            mapping.append(None)

    return mapping


def unique_path(folder, base_name, ext=".xlsx"):
    path = os.path.join(folder, f"{base_name}{ext}")
    if not os.path.exists(path):
        return path
    i = 1
    while True:
        path = os.path.join(folder, f"{base_name} ({i}){ext}")
        if not os.path.exists(path):
            return path
        i += 1


# ============================================================
# СОХРАНЕНИЕ В РАЗНЫЕ ФОРМАТЫ
# ============================================================
def _save_xlsx(out_path, fmt_columns, excel_rows, out_format):
    wb = Workbook()
    ws = wb.active
    ws.title = out_format

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, (name, _) in enumerate(fmt_columns, start=1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = bold
        c.alignment = center
        c.border = border

    for r_off, row_data in enumerate(excel_rows):
        excel_row = 2 + r_off
        for col_idx, value in enumerate(row_data, start=1):
            c = ws.cell(row=excel_row, column=col_idx, value=value)
            c.alignment = left
            c.border = border

    for col_idx in range(1, len(fmt_columns) + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        width = min(max(max_len + 2, 10), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(out_path)


def _save_ods(out_path, fmt_columns, excel_rows):
    data = OrderedDict()
    sheet_data = []

    header = [name for name, _ in fmt_columns]
    sheet_data.append(header)

    for row_data in excel_rows:
        clean_row = ["" if v is None else v for v in row_data]
        sheet_data.append(clean_row)

    data["Sheet1"] = sheet_data
    save_data(out_path, data)


# ============================================================
# ОСНОВНАЯ ФУНКЦИЯ ПЕРЕНОСА
# ============================================================
def do_transfer(word_path, format_key, out_format="xlsx", out_path_manual=None):
    fmt_columns = FORMATS[format_key]

    raw = read_word_tables(word_path)
    if not raw:
        raise ValueError("В Word не найдено ни одной таблицы.")

    idx, w_head, w_data = detect_word_header(raw)

    print("=" * 70)
    print(f"Формат: {format_key}. Определена шапка в строке №{idx + 1} (всего строк: {len(raw)})")
    print("=" * 70)
    print("\nШАПКА WORD:")
    for i, h in enumerate(w_head):
        print(f"  [{i}] {repr(h)}")
    print(f"\nПервые 3 строки данных:")
    for r_i, r in enumerate(w_data[:3]):
        print(f"  --- строка {r_i + 1} (всего ячеек: {len(r)}) ---")
        for c_i, c in enumerate(r):
            print(f"    [{c_i}] {repr(c)}")
    print("=" * 70)

    if not w_data:
        raise ValueError("В таблице Word нет строк с данными.")

    mapping = build_mapping(w_head, fmt_columns)

    print("\nСОПОСТАВЛЕНИЕ КОЛОНОК:")
    for col_idx, ((fmt_name, _), w_idx) in enumerate(zip(fmt_columns, mapping)):
        if w_idx == "AUTO":
            print(f"  ★ [{col_idx}] {fmt_name}  — АВТОЗАПОЛНЕНИЕ")
        elif w_idx is None:
            print(f"  ✗ [{col_idx}] {fmt_name}  — НЕ НАЙДЕНО (пусто)")
        else:
            print(f"  ✓ [{col_idx}] {fmt_name}  ←  Word[{w_idx}] = {repr(w_head[w_idx])}")
    print("=" * 70)

    excel_rows = []
    for wrow in w_data:
        row_data = []
        for w_idx in mapping:
            if w_idx == "AUTO":
                value = format_key
            elif w_idx is None:
                value = None
            else:
                value = wrow[w_idx] if w_idx < len(wrow) else None
                if value == "" or value is None:
                    value = None
            row_data.append(value)
        excel_rows.append(row_data)

    # Определяем куда сохранять
    if out_path_manual:
        out_path = out_path_manual
        folder = os.path.dirname(os.path.abspath(out_path)) or "."
        os.makedirs(folder, exist_ok=True)
    else:
        folder = os.path.dirname(os.path.abspath(word_path))
        ext = EXT_MAP.get(out_format, ".xlsx")
        out_path = unique_path(folder, format_key, ext=ext)

    if out_format == "ods":
        _save_ods(out_path, fmt_columns, excel_rows)
    else:
        _save_xlsx(out_path, fmt_columns, excel_rows, out_format)

    return out_path


# ============================================================
# GUI
# ============================================================
class App:
    def __init__(self, root):
        self.root = root
        root.title("Word → Excel: перенос по формату")
        root.geometry("720x440")
        root.minsize(680, 420)
        root.resizable(False, False)

        self.word_path = tk.StringVar()
        self.format_var = tk.StringVar(value="")
        self.out_format_var = tk.StringVar(value="xlsx")
        self.out_path_var = tk.StringVar()

        notebook = ttk.Notebook(root)
        notebook.pack(fill="both", expand=True, padx=10, pady=10)

        # ============ ВКЛАДКА 1: Основное ============
        tab1 = ttk.Frame(notebook)
        notebook.add(tab1, text="  Основное  ")

        pad = {"padx": 10, "pady": 8}

        tk.Label(tab1, text="Файл Word (.docx):", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", **pad)
        f1 = tk.Frame(tab1)
        f1.pack(fill="x", padx=10)
        tk.Entry(f1, textvariable=self.word_path).pack(side="left", fill="x", expand=True)
        tk.Button(f1, text="Выбрать…", width=12, command=self.pick_word).pack(side="left", padx=5)

        tk.Label(tab1, text="Формат (32 / 35 / 38):", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", padx=10, pady=(14, 4))
        f2 = tk.Frame(tab1)
        f2.pack(fill="x", padx=10)
        self.format_menu = tk.OptionMenu(f2, self.format_var, "", "32", "35", "38")
        self.format_menu.config(width=20, anchor="w")
        self.format_menu.pack(side="left")

        # ============ ВКЛАДКА 2: Дополнительно ============
        tab2 = ttk.Frame(notebook)
        notebook.add(tab2, text="  Дополнительно  ")

        tk.Label(tab2, text="Формат выходного файла:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", padx=10, pady=(10, 4))
        f3 = tk.Frame(tab2)
        f3.pack(fill="x", padx=10)
        self.out_format_menu = tk.OptionMenu(
            f3, self.out_format_var, "xlsx", "xlsm", "ods"
        )
        self.out_format_menu.config(width=20, anchor="w")
        self.out_format_menu.pack(side="left")

        tk.Label(tab2, text="Куда сохранить (пусто — рядом с Word-файлом):",
                 anchor="w", font=("Arial", 11, "bold")).pack(fill="x", padx=10, pady=(18, 4))
        f4 = tk.Frame(tab2)
        f4.pack(fill="x", padx=10)
        tk.Entry(f4, textvariable=self.out_path_var).pack(side="left", fill="x", expand=True)
        tk.Button(f4, text="Выбрать…", width=12, command=self.pick_save_path).pack(side="left", padx=5)

        tk.Label(tab2,
                 text="Форматы: xlsx — новый Excel, xlsm — Excel с макросами,\n"
                      "ods — LibreOffice / OpenOffice",
                 anchor="w", justify="left", fg="#666",
                 font=("Arial", 9)).pack(fill="x", padx=10, pady=(16, 0))

        # ============ Кнопка запуска ============
        self.btn_run = tk.Button(root, text="▶  ЗАПУСТИТЬ",
                                 command=self.run,
                                 bg="#4CAF50", fg="white",
                                 font=("Arial", 13, "bold"),
                                 height=2, cursor="hand2")
        self.btn_run.pack(fill="x", padx=10, pady=10)

    def pick_word(self):
        p = filedialog.askopenfilename(
            title="Выберите Word-файл",
            filetypes=[("Word", "*.docx"), ("Все файлы", "*.*")]
        )
        if p:
            self.word_path.set(p)

    def pick_save_path(self):
        fmt = self.out_format_var.get()
        ext = EXT_MAP.get(fmt, ".xlsx")

        p = filedialog.asksaveasfilename(
            title="Куда сохранить результат",
            defaultextension=ext,
            filetypes=[
                ("Excel 2007+", "*.xlsx"),
                ("Excel с макросами", "*.xlsm"),
                ("OpenDocument", "*.ods"),
                ("Все файлы", "*.*"),
            ]
        )
        if p:
            self.out_path_var.set(p)

    def run(self):
        w = self.word_path.get()
        fmt = self.format_var.get()
        if not w:
            messagebox.showwarning("Не хватает данных", "Укажите Word-файл.")
            return
        if fmt not in FORMATS:
            messagebox.showwarning("Не выбран формат", "Выберите формат: 32, 35 или 38.")
            return

        out_fmt = self.out_format_var.get()
        out_path = self.out_path_var.get().strip() or None

        self.btn_run.config(state="disabled", text="Обработка…")
        threading.Thread(target=self._worker, args=(w, fmt, out_fmt, out_path), daemon=True).start()

    def _worker(self, w, fmt, out_fmt, out_path):
        try:
            result = do_transfer(w, fmt, out_format=out_fmt, out_path_manual=out_path)
            self.btn_run.config(state="normal", text="▶  ЗАПУСТИТЬ")
            messagebox.showinfo("Готово", f"Файл сохранён:\n{result}")
        except Exception as ex:
            self.btn_run.config(state="normal", text="▶  ЗАПУСТИТЬ")
            messagebox.showerror("Ошибка", f"{ex}\n\n{traceback.format_exc()}")


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
